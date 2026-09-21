# -*- coding: utf-8 -*-
"""從 投資計算表.xlsx 依標籤定位抽出各物件資料（不寫死列號）"""
import openpyxl, json, sys

SRC = r"C:\Users\Joan\Desktop\投資計算表.xlsx"
SKIP = {"步驟", "提供項目", "空白表單"}

def find(ws, *labels, col=(1, 3), maxr=60):
    """回傳第一個符合標籤的 (row, col)"""
    for r in range(1, maxr):
        for c in col:
            v = ws.cell(r, c).value
            if isinstance(v, str):
                s = v.strip().replace("\n", "")
                for lb in labels:
                    if s.startswith(lb):
                        return r, c
    return None, None

def val(ws, *labels, off=1, col=(1, 3)):
    r, c = find(ws, *labels, col=col)
    if r is None:
        return None
    return ws.cell(r, c + off).value

def num(v):
    return v if isinstance(v, (int, float)) else None

def extract(ws):
    d = {"sheet": ws.title}
    d["addr"] = val(ws, "地址")
    d["ping"] = num(val(ws, "總坪數"))
    d["type"] = val(ws, "房屋型態", col=(3,))
    d["age"] = num(val(ws, "建築完成日期", col=(3,)))
    d["price"] = num(val(ws, "購入價格"))
    d["appr"] = num(val(ws, "銀行鑑價"))
    d["apprAvg"] = num(val(ws, "銀行估價平均"))
    d["rentTotal"] = num(val(ws, "月總租金"))
    d["roi"] = num(val(ws, "ROI"))
    d["yield"] = num(val(ws, "總價投報率"))
    d["cashflow"] = num(val(ws, "月現金流"))

    # 三家銀行估價：在「銀行估價平均」上方找
    r15, _ = find(ws, "銀行估價平均")
    banks = []
    if r15:
        for r in range(max(1, r15 - 4), r15):
            nm, v = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(nm, str) and isinstance(v, (int, float)) and nm.strip() not in ("銀行",):
                banks.append({"name": nm.strip(), "appr": v})
    d["banks"] = banks

    # 貸款
    r17, _ = find(ws, "一順位")
    if r17:
        d["loan1"] = num(ws.cell(r17, 2).value)
        d["rate1"] = num(ws.cell(r17 + 1, 2).value)
        d["loan2"] = num(ws.cell(r17, 4).value)
        d["rate2"] = num(ws.cell(r17 + 1, 4).value)

    # 租金四戶
    rr, _ = find(ws, "租金", col=(1,))
    rents = []
    if rr:
        for r in range(rr + 1, rr + 5):
            if ws.cell(r, 1).value in ("A", "B", "C", "D"):
                rents.append(num(ws.cell(r, 2).value))
    d["rents"] = rents

    # 成本明細：找「成本計算」下方 C/D 欄
    rc, cc = find(ws, "成本計算", col=(3,))
    cost = {}
    if rc:
        for r in range(rc + 1, rc + 12):
            k, v = ws.cell(r, 3).value, ws.cell(r, 4).value
            if isinstance(k, str) and isinstance(v, (int, float)):
                k = k.strip()
                if k.startswith("總計"):
                    d["costTotal"] = v
                    break
                cost[k] = v
    d["cost"] = cost
    d["renov"] = num(val(ws, "裝修總計", col=(3,)))
    return d

def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.title in SKIP:
            continue
        d = extract(ws)
        if d.get("addr") and d.get("price"):
            out.append(d)
    json.dump(out, open("cases.json", "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(f"抽出 {len(out)} 個物件：")
    for d in out:
        print(f"  {d['sheet']:<14} {d['addr'][:28]:<30} 購入 {d['price']}萬 "
              f"估價 {d.get('apprAvg')} 租金 {d.get('rentTotal')}")

if __name__ == "__main__":
    main()
